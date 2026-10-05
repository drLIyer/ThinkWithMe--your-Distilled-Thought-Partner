"""
Distill — Chainlit UI

Run with:
    chainlit run chainlit_app.py
"""

import asyncio
import base64
import io
import json
import mimetypes
import os
import re
import signal
import sys
import time
import uuid
from datetime import datetime
from pathlib import Path

from dotenv import load_dotenv
load_dotenv()

import anthropic
import httpx
import numpy as np

import chainlit as cl
from chainlit.data.sql_alchemy import SQLAlchemyDataLayer
from chainlit.data.storage_clients.base import BaseStorageClient
from chainlit.user import User


class _NoOpStorage(BaseStorageClient):
    """Stub storage client — satisfies Chainlit's interface without persisting files."""
    async def upload_file(self, object_key, data, mime="application/octet-stream", overwrite=True, content_disposition=None):
        return {"object_key": object_key}  # non-empty so Chainlit's truthiness check passes
    async def delete_file(self, object_key):
        return True
    async def get_read_url(self, object_key):
        return ""
    async def close(self):
        pass

_BASE              = Path(__file__).parent
DB_URL             = f"sqlite+aiosqlite:///{_BASE / 'asklenny.db'}"
DB_SYNC_URL        = str(_BASE / "asklenny.db")
USER_PROFILE_PATH  = _BASE / "user_profile.json"
USER_MEMORIES_PATH = _BASE / "user_memories.json"
BOOKMARKS_PATH     = _BASE / "bookmarks.json"


@cl.data_layer
def get_data_layer():
    return SQLAlchemyDataLayer(DB_URL, storage_provider=_NoOpStorage())


@cl.password_auth_callback
async def auth_callback(username: str, password: str):
    # Simple single-user auth — same identifier as migration
    if username == os.environ.get("CHAINLIT_USERNAME", "") and password == os.environ.get("CHAINLIT_PASSWORD", ""):
        return User(identifier="local-user", metadata={"role": "user"})
    return None

@cl.on_settings_update
async def on_settings_update(settings: dict):
    # If a bookmark was selected, open it in the chat
    bookmark_id = settings.get("_open_bookmark", "")
    if bookmark_id and bookmark_id != "__none__":
        bookmarks = await _get_bookmarks_list()
        match = next((b for b in bookmarks if b["id"] == bookmark_id), None)
        if match:
            import re as _re
            sources_md = match.get("sources_md", "")
            # Extract all source links for display
            src_blocks = _re.findall(
                r'(🎙️|📰) \*\*(Podcast|Newsletter)\*\* · ([^·]+) · \*[^*]+\* · `[^`]*`([^\n]*)',
                sources_md
            )
            src_lines = []
            for icon, kind, label, link_part in src_blocks:
                label = label.strip()
                open_match = _re.search(r'\[Open ↗\]\((https?://[^)]+)\)', link_part)
                yt_match   = _re.search(r'\[Search on YouTube ↗\]\((https?://[^)]+)\)', link_part)
                url = open_match or yt_match
                src_lines.append(f"{icon} [{label}]({url.group(1)})" if url else f"{icon} {label}")

            date = match.get("timestamp", "")[:10]
            content = (
                f"### 📌 {match['question']}\n"
                f"*Saved on {date}*\n\n"
                f"{match['answer']}"
            )
            if src_lines:
                content += "\n\n**Sources**\n" + "  \n".join(src_lines)

            thread_id = cl.user_session.get("thread_id", "")
            asyncio.ensure_future(_post_engagement("bookmark_open", match["question"], thread_id))
            await cl.Message(content=content, author="Distill").send()
        return

    # Otherwise save profile as normal
    profile = {k: settings.get(k, "") for k in ("role", "team", "products", "projects", "goals")}
    ok = await _put_user_profile(profile)
    if not ok:
        save_user_profile(profile)
    cl.user_session.set("user_profile", profile)

    await cl.Message(
        content=f"✅ Settings saved — I'll tailor answers to your role as **{profile.get('role') or 'your role'}** on **{profile.get('team') or 'your team'}**.",
        author="Distill",
    ).send()


# ── Config ────────────────────────────────────────────────────────────────────

FEEDBACK_PATH      = _BASE / "feedback.json"
CLAUDE_MODEL       = os.environ.get("ANTHROPIC_MODEL",       "claude-sonnet-4-6")
HAIKU_MODEL        = os.environ.get("ANTHROPIC_HAIKU_MODEL", "claude-haiku-4-5")
MAX_HISTORY_TURNS    = 20
TURN_WARNING         = 15
MAX_MEMORIES         = 10
MAX_UPLOADED_DOCS    = 3
SIMILAR_Q_THRESHOLD  = 0.82
RAG_SERVICE_URL           = "http://localhost:8083"
RAG_SERVICE_TIMEOUT       = 10.0
BOOKMARKS_SERVICE_URL     = "http://localhost:8084"
BOOKMARKS_SERVICE_TIMEOUT = 5.0
USERS_SERVICE_URL         = "http://localhost:8085"
USERS_SERVICE_TIMEOUT     = 5.0
FEEDBACK_SERVICE_URL      = "http://localhost:8086"
FEEDBACK_SERVICE_TIMEOUT  = 5.0

_ASSISTANT_NAME   = os.environ.get("ASSISTANT_NAME",        "Distill")
_CORPUS_DESC      = os.environ.get("CORPUS_DESCRIPTION",
    "Lenny Rachitsky's newsletter and podcast: 600+ pieces on product, growth, leadership, and startups")
_CUSTOM_PROMPT    = os.environ.get("SYSTEM_PROMPT_OVERRIDE", "")

SYSTEM_PROMPT = _CUSTOM_PROMPT or f"""\
You are {_ASSISTANT_NAME} — a thought partner with deep knowledge of {_CORPUS_DESC}.

Always cite sources naturally in the text (e.g. "In the podcast with [Guest]..." or "In the newsletter '[Title]'...").

When web search results are provided under "## Web Search Results", use them to supplement your answer for topics the corpus hasn't covered directly. Cite web sources inline as [Source: title](url).

Adapt your response to the user's intent:

**Direct questions**: Answer concisely using the provided context. If the context is thin, say so — but share what is relevant.

**Draft or writing requests** ("draft an article", "write a post", "help me write X"):
Use the context to produce a concrete draft the user can act on immediately, even if no single source covers the topic perfectly. Present it under the heading "Here's a version you could start with:" and note which sources shaped it.

**How-to or brainstorm requests** ("how should I approach", "brainstorm ideas", "help me think through"):
1. Present a brief framework grounded in the sources.
2. Offer 2–3 specific directions or angles the user could take, each as a numbered option.
3. End with: "Which of these would you like to go deeper on?"

**Always end every response** with a section formatted exactly like this:

---
**Key Takeaways**
- [2–3 bullet takeaways from this response]

*As your thought partner — what would you like to explore next? [Ask one specific, relevant follow-up question]*

**Source transparency**: Always be honest about where your answer comes from.
- If the provided context strongly supports your answer, cite it naturally.
- If the context is thin or only partially relevant, say so explicitly.
- Never silently blend general knowledge into a sourced answer without flagging it.

Never leave the user without something concrete to act on."""

FOLLOWUP_PROMPT_TEMPLATE = """\
Based on this conversation exchange, suggest exactly 3 short follow-up questions the user might want to ask next.
The user is {level_desc}. Calibrate the questions to their level:
- Executive/VP: strategic, org-level, outcome-focused
- Director: balance strategy with team execution
- Senior PM/IC: craft, judgment, stakeholder influence
- PM: practical, concrete, how-to
Return ONLY a JSON array of 3 strings, no other text. Each question should be under 10 words."""


def _detect_seniority_local(role: str) -> str:
    r = role.lower()
    if any(t in r for t in ["svp", "vp ", "vp,", "chief", "president", "head of", "evp", "gm "]):
        return "executive"
    if any(t in r for t in ["director", "sr. director", "senior director"]):
        return "director"
    if any(t in r for t in ["senior pm", "sr. pm", "principal", "lead pm", "staff pm", "group pm"]):
        return "senior_ic"
    if any(t in r for t in ["product manager", " pm", "pm,", "pm "]):
        return "pm"
    return "unknown"

def _graceful_shutdown(signum, frame):
    """Clean up multiprocessing resources before exit to prevent semaphore leaks."""
    import multiprocessing
    for child in multiprocessing.active_children():
        child.terminate()
        child.join(timeout=2)
    sys.exit(0)

signal.signal(signal.SIGTERM, _graceful_shutdown)
signal.signal(signal.SIGINT, _graceful_shutdown)


async def _call_rag_service(query: str, history: list[dict] | None = None) -> dict | None:
    """Call POST /retrieve/with-context on the RAG service. Returns None on any failure."""
    try:
        async with httpx.AsyncClient(timeout=RAG_SERVICE_TIMEOUT) as client:
            resp = await client.post(
                f"{RAG_SERVICE_URL}/retrieve/with-context",
                json={"query": query, "history": history or []},
            )
            resp.raise_for_status()
            return resp.json()
    except Exception:
        return None


async def _get_bookmarks_render() -> str | None:
    """Returns rendered markdown from bookmarks service, or None on failure."""
    try:
        async with httpx.AsyncClient(timeout=BOOKMARKS_SERVICE_TIMEOUT) as client:
            resp = await client.get(f"{BOOKMARKS_SERVICE_URL}/bookmarks/render")
            resp.raise_for_status()
            return resp.json()["markdown"]
    except Exception:
        return None


async def _post_bookmark(question: str, answer: str, sources_md: str,
                         thread_id: str, thread_name: str) -> bool:
    """Posts a new bookmark to the bookmarks service. Returns True on success."""
    try:
        async with httpx.AsyncClient(timeout=BOOKMARKS_SERVICE_TIMEOUT) as client:
            resp = await client.post(
                f"{BOOKMARKS_SERVICE_URL}/bookmarks",
                json={"question": question, "answer": answer, "sources_md": sources_md,
                      "thread_id": thread_id, "thread_name": thread_name},
            )
            resp.raise_for_status()
            return True
    except Exception:
        return False


async def _get_system_prompt() -> str | None:
    """Returns full system prompt (with profile+memories prefix) from users service."""
    try:
        async with httpx.AsyncClient(timeout=USERS_SERVICE_TIMEOUT) as client:
            resp = await client.get(f"{USERS_SERVICE_URL}/system-prompt")
            resp.raise_for_status()
            return resp.json()["system_prompt"]
    except Exception:
        return None


async def _get_user_profile() -> dict | None:
    """Returns user profile dict from users service."""
    try:
        async with httpx.AsyncClient(timeout=USERS_SERVICE_TIMEOUT) as client:
            resp = await client.get(f"{USERS_SERVICE_URL}/profile")
            resp.raise_for_status()
            return resp.json()
    except Exception:
        return None


async def _put_user_profile(profile: dict) -> bool:
    """Saves user profile to users service. Returns True on success."""
    try:
        async with httpx.AsyncClient(timeout=USERS_SERVICE_TIMEOUT) as client:
            resp = await client.put(f"{USERS_SERVICE_URL}/profile", json=profile)
            resp.raise_for_status()
            return True
    except Exception:
        return False


async def _post_memory(query: str, response_text: str, thread_id: str) -> bool:
    """Sends exchange to users service to extract and save a memory summary."""
    try:
        async with httpx.AsyncClient(timeout=30.0) as client:
            resp = await client.post(
                f"{USERS_SERVICE_URL}/memories",
                json={"query": query, "response_text": response_text, "thread_id": thread_id},
            )
            resp.raise_for_status()
            return True
    except Exception:
        return False


async def _post_engagement(event_type: str, content: str, thread_id: str) -> None:
    """Fire-and-forget engagement event logging to users service."""
    try:
        async with httpx.AsyncClient(timeout=USERS_SERVICE_TIMEOUT) as client:
            await client.post(
                f"{USERS_SERVICE_URL}/engagement",
                json={"type": event_type, "content": content, "thread_id": thread_id},
            )
    except Exception:
        pass


async def _post_feedback(msg_id: str, question: str, answer: str,
                         rating: str, feedback_text: str = "") -> bool:
    """Posts feedback entry to the feedback service. Returns True on success."""
    try:
        async with httpx.AsyncClient(timeout=FEEDBACK_SERVICE_TIMEOUT) as client:
            resp = await client.post(
                f"{FEEDBACK_SERVICE_URL}/feedback",
                json={"id": msg_id, "question": question, "answer": answer,
                      "rating": rating, "feedback_text": feedback_text},
            )
            resp.raise_for_status()
            return True
    except Exception:
        return False


async def _post_gap(query: str, confidence: dict, sources: list) -> bool:
    """Posts a knowledge gap entry to the feedback service. Returns True on success."""
    try:
        top_score   = round(sources[0]["score"], 4) if sources else 0
        avg_score   = round(sum(s["score"] for s in sources) / len(sources), 4) if sources else 0
        best_source = sources[0].get("title", "") if sources else ""
        async with httpx.AsyncClient(timeout=FEEDBACK_SERVICE_TIMEOUT) as client:
            resp = await client.post(
                f"{FEEDBACK_SERVICE_URL}/gaps",
                json={"query": query, "top_score": top_score, "avg_score": avg_score,
                      "best_source": best_source, "confidence": confidence.get("label", "")},
            )
            resp.raise_for_status()
            return True
    except Exception:
        return False


# ── Attachment handling ───────────────────────────────────────────────────────

SUPPORTED_IMAGE_TYPES = {"image/jpeg", "image/png", "image/gif", "image/webp"}


def process_attachment(file_element) -> dict:
    name = file_element.name
    raw  = getattr(file_element, "content", None)
    if not raw and getattr(file_element, "path", None):
        with open(file_element.path, "rb") as fh:
            raw = fh.read()
    mime, _ = mimetypes.guess_type(name)
    mime = mime or "application/octet-stream"

    if mime in SUPPORTED_IMAGE_TYPES:
        return {"kind": "image", "name": name, "mime": mime,
                "data": base64.b64encode(raw).decode()}

    if mime in {"text/plain", "text/markdown", "text/csv", "application/json"} \
            or name.endswith((".md", ".txt", ".csv")):
        return {"kind": "text", "name": name, "mime": mime,
                "content": raw.decode("utf-8", errors="replace")}

    if mime == "application/pdf" or name.lower().endswith(".pdf"):
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(raw))
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        except ImportError:
            text = "[PDF uploaded but pypdf is not installed — run: pip install pypdf]"
        return {"kind": "text", "name": name, "mime": "application/pdf", "content": text}

    if mime in {
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "application/vnd.ms-powerpoint",
    } or name.lower().endswith((".pptx", ".ppt")):
        if name.lower().endswith(".ppt") and not name.lower().endswith(".pptx"):
            return {"kind": "text", "name": name, "mime": mime,
                    "content": "[Legacy .ppt format is not supported — please save as .pptx and re-upload]"}
        try:
            from pptx import Presentation
            prs = Presentation(io.BytesIO(raw))
            slides_text = []
            for i, slide in enumerate(prs.slides, 1):
                texts = [
                    shape.text.strip()
                    for shape in slide.shapes
                    if hasattr(shape, "text") and shape.text.strip()
                ]
                if texts:
                    slides_text.append(f"--- Slide {i} ---\n" + "\n".join(texts))
            text = "\n\n".join(slides_text) if slides_text else "[No text found in presentation]"
        except ImportError:
            text = "[PowerPoint uploaded but python-pptx is not installed — run: pip install python-pptx]"
        except Exception as e:
            text = f"[Could not read PowerPoint file: {e}]"
        return {"kind": "text", "name": name, "mime": mime, "content": text}

    if mime in {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    } or name.lower().endswith((".docx", ".doc")):
        if name.lower().endswith(".doc") and not name.lower().endswith(".docx"):
            return {"kind": "text", "name": name, "mime": mime,
                    "content": "[Legacy .doc format is not supported — please save as .docx and re-upload]"}
        try:
            from docx import Document
            doc = Document(io.BytesIO(raw))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            text = "\n\n".join(paragraphs) if paragraphs else "[No text found in document]"
        except ImportError:
            text = "[Word document uploaded but python-docx is not installed — run: pip install python-docx]"
        except Exception as e:
            text = f"[Could not read Word document: {e}]"
        return {"kind": "text", "name": name, "mime": mime, "content": text}

    try:
        content = raw.decode("utf-8", errors="replace")
    except Exception:
        content = "[Binary file — cannot display content]"
    return {"kind": "text", "name": name, "mime": mime, "content": content}


def attachments_to_claude_blocks(attachments: list[dict], query: str) -> list:
    blocks = []
    for att in attachments:
        if att["kind"] == "image":
            blocks.append({"type": "image", "source": {
                "type": "base64", "media_type": att["mime"], "data": att["data"]}})
        else:
            blocks.append({"type": "text",
                           "text": f"[Attached file: {att['name']}]\n\n{att['content']}"})
    blocks.append({"type": "text", "text": query})
    return blocks


# ── Feedback ─────────────────────────────────────────────────────────────────

async def save_feedback(msg_id: str, question: str, answer: str, rating: str, text: str = ""):
    # Write via feedback service (file-locked); fallback to direct write if service is down
    ok = await _post_feedback(msg_id, question, answer, rating, text)
    if not ok:
        records = []
        if FEEDBACK_PATH.exists():
            try:
                records = json.loads(FEEDBACK_PATH.read_text())
            except Exception:
                records = []
        records.append({
            "id": msg_id,
            "timestamp": datetime.utcnow().isoformat(),
            "question": question,
            "answer": answer,
            "rating": rating,
            "feedback_text": text,
        })
        FEEDBACK_PATH.write_text(json.dumps(records, indent=2))

    # Also write to SQLite
    try:
        import aiosqlite
        value = 1 if rating == "up" else 0
        thread_id = cl.user_session.get("thread_id", "unknown")
        async with aiosqlite.connect(DB_SYNC_URL) as db:
            await db.execute(
                'INSERT OR REPLACE INTO feedbacks (id, "forId", "threadId", value, comment) VALUES (?,?,?,?,?)',
                (str(uuid.uuid4()), msg_id, thread_id, value,
                 f"Q: {question[:200]}\n\nA: {answer[:500]}\n\nFeedback: {text}"),
            )
            await db.commit()
    except Exception:
        pass


# ── Follow-up suggestions ─────────────────────────────────────────────────────

def get_followup_suggestions(question: str, answer: str, role: str = "") -> list[str]:
    level = _detect_seniority_local(role)
    level_desc = {
        "executive": "a VP/SVP/executive leader",
        "director":  "a Director",
        "senior_ic": "a Senior PM or Lead PM",
        "pm":        "a Product Manager",
        "unknown":   "a product professional",
    }.get(level, "a product professional")
    prompt = FOLLOWUP_PROMPT_TEMPLATE.format(level_desc=level_desc)
    try:
        client = anthropic.Anthropic()
        resp = client.messages.create(
            model=CLAUDE_MODEL,
            max_tokens=200,
            messages=[{"role": "user",
                       "content": f"{prompt}\n\nQ: {question}\nA: {answer[:600]}"}],
        )
        suggestions = json.loads(resp.content[0].text.strip())
        if isinstance(suggestions, list):
            return suggestions[:3]
    except Exception:
        pass
    return []


# ── User profile ─────────────────────────────────────────────────────────────

def load_user_profile() -> dict:
    if USER_PROFILE_PATH.exists():
        try:
            return json.loads(USER_PROFILE_PATH.read_text())
        except Exception:
            pass
    return {}


def save_user_profile(profile: dict):
    profile["updated_at"] = datetime.utcnow().isoformat()
    USER_PROFILE_PATH.write_text(json.dumps(profile, indent=2))


def _build_system_prompt() -> str:
    profile = cl.user_session.get("user_profile", {})
    memories = cl.user_session.get("memories", [])

    prefix = ""

    if any(v for k, v in profile.items() if k != "updated_at" and v):
        goals = profile.get("goals", "")
        prefix += (
            "## About this user\n"
            f"- Role: {profile.get('role', '')}\n"
            f"- Team / Business unit: {profile.get('team', '')}\n"
            f"- Products / Portfolios: {profile.get('products', '')}\n"
            f"- Current projects: {profile.get('projects', '')}\n"
            + (f"- What they want from this thought partner: {goals}\n" if goals else "")
            + "\nAlways frame answers in the context of their specific role and the products they manage. "
            "When relevant, connect frameworks from the corpus to their actual situation."
            + (f" Keep in mind their stated goal: {goals}." if goals else "")
            + "\n\n"
        )

    if memories:
        prefix += "## Context from your recent sessions\n"
        prefix += "\n".join(f"- {m['summary']}" for m in memories[-MAX_MEMORIES:])
        prefix += "\n\n"

    return (prefix + SYSTEM_PROMPT) if prefix else SYSTEM_PROMPT


async def _get_bookmarks_list() -> list:
    """Returns list of bookmark dicts from the bookmarks service, or local fallback."""
    try:
        async with httpx.AsyncClient(timeout=BOOKMARKS_SERVICE_TIMEOUT) as client:
            resp = await client.get(f"{BOOKMARKS_SERVICE_URL}/bookmarks")
            resp.raise_for_status()
            return resp.json().get("bookmarks", [])
    except Exception:
        return load_bookmarks()


async def _send_chat_settings():
    profile   = cl.user_session.get("user_profile", {})
    bookmarks = await _get_bookmarks_list()

    # Build {truncated question: bookmark id} for the Select widget
    # Most recent 20, newest first
    bookmark_items = {
        b["question"][:80]: b["id"]
        for b in reversed(bookmarks[-20:])
    } if bookmarks else {"(no saved posts yet)": "__none__"}

    settings = cl.ChatSettings([
        cl.input_widget.Tab(id="profile_tab", label="Profile", inputs=[
            cl.input_widget.TextInput(id="role",     label="Your role",
                                      placeholder="e.g. Sr. PM, Director of Product",
                                      initial=profile.get("role", "")),
            cl.input_widget.TextInput(id="team",     label="Team / Business unit",
                                      placeholder="e.g. Orthodontics, iTero, Scanner",
                                      initial=profile.get("team", "")),
            cl.input_widget.TextInput(id="products", label="Products or portfolios you manage",
                                      placeholder="e.g. iTero Element, ClinCheck, Vivera",
                                      initial=profile.get("products", "")),
            cl.input_widget.TextInput(id="projects", label="Current projects or initiatives",
                                      placeholder="e.g. Q3 launch, scanner pilots, growth initiative",
                                      initial=profile.get("projects", "")),
            cl.input_widget.TextInput(id="goals",    label="What do you most want from your thought partner?",
                                      placeholder="e.g. Help me think strategically, challenge my assumptions, help me write better",
                                      initial=profile.get("goals", "")),
        ]),
        cl.input_widget.Tab(id="saved_posts_tab", label="📌 Saved Posts", inputs=[
            cl.input_widget.Select(
                id="_open_bookmark",
                label="Select a saved post to open it in the chat",
                items=bookmark_items,
            ),
        ]),
    ])
    await settings.send()


# ── Topic memory ──────────────────────────────────────────────────────────────

def load_memories() -> list:
    if USER_MEMORIES_PATH.exists():
        try:
            return json.loads(USER_MEMORIES_PATH.read_text())
        except Exception:
            pass
    return []


# ── Bookmarks ─────────────────────────────────────────────────────────────────

def load_bookmarks() -> list:
    if BOOKMARKS_PATH.exists():
        try:
            return json.loads(BOOKMARKS_PATH.read_text())
        except Exception:
            pass
    return []


def save_bookmark(question: str, answer: str, sources_md: str, thread_id: str, thread_name: str):
    bookmarks = load_bookmarks()
    bookmarks.append({
        "id": str(uuid.uuid4()),
        "timestamp": datetime.utcnow().isoformat(),
        "thread_id": thread_id,
        "thread_name": thread_name or "Untitled",
        "question": question,
        "answer": answer,
        "sources_md": sources_md,
    })
    BOOKMARKS_PATH.write_text(json.dumps(bookmarks, indent=2))


def _render_bookmarks_md() -> str:
    bookmarks = load_bookmarks()
    if not bookmarks:
        return "_No saved posts yet. Click 📌 on any answer to save it here._"
    lines = []
    for b in reversed(bookmarks[-20:]):
        date = b.get("timestamp", "")[:10]
        q = b.get("question", "")[:100]
        a = b.get("answer", "")[:300]
        lines.append(f"**{q}**  \n*{date}*  \n{a}…\n\n---")
    return "\n\n".join(lines)


# ── Web search ───────────────────────────────────────────────────────────────

_WEB_SEARCH_ENGINE = os.environ.get("WEB_SEARCH_ENGINE", "ddgs").lower()


def _web_search_sync(query: str, max_results: int = 5) -> list[dict]:
    """Returns list of {title, href, body} dicts. Empty list on any failure or disabled."""
    if _WEB_SEARCH_ENGINE == "none":
        return []
    try:
        if _WEB_SEARCH_ENGINE == "brave":
            return _web_search_brave(query, max_results)
        return _web_search_ddgs(query, max_results)
    except Exception:
        return []


def _web_search_ddgs(query: str, max_results: int) -> list[dict]:
    from ddgs import DDGS
    raw = DDGS().text(query, max_results=max_results * 2)
    results = [r for r in raw if r.get("href") and "aclick" not in r["href"] and "doubleclick" not in r["href"]]
    return results[:max_results]


def _web_search_brave(query: str, max_results: int) -> list[dict]:
    key = os.environ.get("BRAVE_SEARCH_API_KEY", "")
    if not key:
        return []
    import httpx as _httpx
    resp = _httpx.get(
        "https://api.search.brave.com/res/v1/web/search",
        params={"q": query, "count": max_results},
        headers={"Accept": "application/json", "X-Subscription-Token": key},
        timeout=10.0,
    )
    resp.raise_for_status()
    return [
        {"title": r.get("title", ""), "href": r.get("url", ""), "body": r.get("description", "")}
        for r in resp.json().get("web", {}).get("results", [])
    ]


def _format_web_context(results: list[dict]) -> str:
    parts = []
    for r in results:
        parts.append(f"**{r.get('title', '')}**\n{r.get('body', '')}\nURL: {r.get('href', '')}")
    return "\n\n".join(parts)


def _format_web_sources_md(results: list[dict]) -> str:
    lines = []
    for r in results:
        title = r.get("title", "Web result")
        url   = r.get("href", "")
        body  = r.get("body", "")[:200].rsplit(" ", 1)[0] + " …"
        lines.append(f"🌐 **Web** · [{title}]({url})\n> {body}\n")
    return "\n".join(lines)


# ── Similar past question ─────────────────────────────────────────────────────

def _find_similar_past_question_sync(query: str, current_thread_id: str) -> dict | None:
    try:
        import sqlite3 as _sq
        conn = _sq.connect(DB_SYNC_URL)
        rows = conn.execute(
            "SELECT s.output, s2.output, t.name, s.createdAt "
            "FROM steps s "
            "JOIN steps s2 ON s2.threadId = s.threadId "
            "JOIN threads t ON t.id = s.threadId "
            "WHERE s.type='user_message' AND s2.type='assistant_message' "
            "AND s.threadId != ? "
            "AND s.output != '' "
            "ORDER BY s.createdAt DESC LIMIT 100",
            (current_thread_id,)
        ).fetchall()
        conn.close()
        if not rows:
            return None

        texts = [query] + [r[0] for r in rows]
        resp = httpx.post(
            f"{RAG_SERVICE_URL}/embed",
            json={"texts": texts},
            timeout=10.0,
        )
        resp.raise_for_status()
        embs   = np.array(resp.json()["embeddings"], dtype=np.float32)
        q_emb  = embs[0:1]
        p_embs = embs[1:]
        scores = (p_embs @ q_emb.T).flatten()
        best   = int(scores.argmax())
        if scores[best] >= SIMILAR_Q_THRESHOLD:
            return {
                "question": rows[best][0],
                "answer":   rows[best][1][:400],
                "thread":   rows[best][2] or "a previous chat",
                "date":     rows[best][3][:10],
            }
    except Exception:
        pass
    return None


# ── Turn helpers ──────────────────────────────────────────────────────────────

def count_turns() -> int:
    history = cl.user_session.get("history", [])
    return sum(1 for m in history if m["role"] == "assistant")


def active_model() -> str:
    return HAIKU_MODEL if count_turns() >= TURN_WARNING else CLAUDE_MODEL


# ── Chainlit lifecycle ────────────────────────────────────────────────────────

@cl.on_chat_start
async def on_chat_start():
    cl.user_session.set("history", [])
    cl.user_session.set("msg_id_map", {})
    cl.user_session.set("uploaded_docs", [])
    cl.user_session.set("thread_id", cl.context.session.thread_id or str(uuid.uuid4()))

    profile = await _get_user_profile() or load_user_profile()
    cl.user_session.set("user_profile", profile)
    cl.user_session.set("memories", load_memories())

    is_first_time = not any(v for k, v in profile.items() if k != "updated_at" and v)

    if is_first_time:
        await cl.Message(
            content=(
                "## Welcome to Distill 👋\n\n"
                "Before we start, tell me a bit about you so every answer is relevant to **your** work.\n\n"
                "Fill in your profile using the **⚙️ Settings** panel that just opened — takes 30 seconds. "
                "You can update it anytime from ⚙️ Settings.\n\n"
                "---\n\n"
                f"*{_CORPUS_DESC}, personalized to your role and products.*"
            ),
            author="Distill",
        ).send()
    else:
        await cl.Message(
            content=(
                "## Distill\n"
                "#### Your thought partner for product, growth, and leadership\n\n"
                f"*{_CORPUS_DESC}*\n\n"
                "---\n\n"
                f"Ask me anything — or attach a file to analyze through the corpus."
            ),
            author="Distill",
        ).send()

    await _send_chat_settings()


@cl.on_chat_resume
async def on_chat_resume(thread: dict):
    raw = []
    for step in thread.get("steps", []):
        stype = step.get("type", "")
        content = step.get("output", "").strip()
        if not content:
            continue
        if stype == "user_message":
            raw.append({"role": "user", "content": content})
        elif stype == "assistant_message":
            # Skip the welcome banner and follow-up suggestion messages
            if content.startswith("## Distill") or content.startswith("**Suggested follow-ups"):
                continue
            # Strip sources/confidence suffix before storing in Claude history
            if "\n\n---\n" in content:
                content = content[:content.index("\n\n---\n")]
            if content.strip():
                raw.append({"role": "assistant", "content": content.strip()})

    # Ensure strictly alternating user/assistant (Claude API requirement).
    # Keep only turns where role alternates, starting from the first user message.
    history = []
    for msg in raw:
        if not history:
            if msg["role"] == "user":
                history.append(msg)
        elif msg["role"] != history[-1]["role"]:
            history.append(msg)
        # Same role in a row — skip to avoid API errors

    if len(history) > MAX_HISTORY_TURNS * 2:
        history = history[-(MAX_HISTORY_TURNS * 2):]

    cl.user_session.set("history", history)
    cl.user_session.set("msg_id_map", {})
    cl.user_session.set("uploaded_docs", [])
    cl.user_session.set("thread_id", thread.get("id", str(uuid.uuid4())))
    profile = await _get_user_profile() or load_user_profile()
    cl.user_session.set("user_profile", profile)
    cl.user_session.set("memories", load_memories())
    await _send_chat_settings()


@cl.on_message
async def on_message(message: cl.Message):
    try:
        await _handle_message(message)
    except Exception as e:
        import traceback
        await cl.Message(content=f"❌ Unexpected error: {e}\n```\n{traceback.format_exc()[:500]}\n```", author="System").send()


async def _handle_message(message: cl.Message):
    history: list[dict] = cl.user_session.get("history", [])
    turns = count_turns()
    model = active_model()

    # ── Turn warnings ──────────────────────────────────────────────────────
    if turns == TURN_WARNING:
        await cl.Message(
            content=(
                f"⚠️ **You've reached {TURN_WARNING} turns.** Long conversations reduce response "
                f"quality as earlier context gets deprioritized.\n\n"
                f"**Recommendation:** Start a new chat (click ✏️ New Chat) to keep answers sharp.\n\n"
                f"If you continue, responses will switch to **Claude Haiku** "
                f"(`{HAIKU_MODEL}`) which is faster but has limitations:\n"
                f"- Less nuanced reasoning\n"
                f"- May miss subtle connections across the corpus\n"
                f"- Shorter, less detailed responses\n"
                f"- Smaller effective context window"
            ),
            author="System",
        ).send()
    elif turns > TURN_WARNING:
        await cl.Message(
            content=f"ℹ️ Using **Claude Haiku** (`{HAIKU_MODEL}`) — start a new chat for full Sonnet quality.",
            author="System",
        ).send()

    # ── Process attachments ────────────────────────────────────────────────
    attachments = []
    uploaded_docs: list = cl.user_session.get("uploaded_docs", [])
    if message.elements:
        for el in message.elements:
            if getattr(el, "content", None) or getattr(el, "path", None):
                att = process_attachment(el)
                attachments.append(att)
                if att["kind"] == "text":
                    uploaded_docs = [d for d in uploaded_docs if d["name"] != att["name"]]
                    uploaded_docs.append({"name": att["name"], "content": att["content"]})
                    if len(uploaded_docs) > MAX_UPLOADED_DOCS:
                        uploaded_docs = uploaded_docs[-MAX_UPLOADED_DOCS:]
                    cl.user_session.set("uploaded_docs", uploaded_docs)
                    await cl.Message(
                        content=f"📄 **{att['name']}** added to context — I'll reference it alongside the corpus.",
                        author="Distill", parent_id=None,
                    ).send()

    # ── Retrieve ───────────────────────────────────────────────────────────
    query = message.content
    thread_id = cl.user_session.get("thread_id", "")

    # Call RAG service and similar-question lookup concurrently
    loop = asyncio.get_running_loop()
    rag_future     = _call_rag_service(query, history)
    similar_future = loop.run_in_executor(None, _find_similar_past_question_sync, query, thread_id)
    rag_response, similar = await asyncio.gather(rag_future, similar_future)

    if rag_response is None:
        await cl.Message(
            content="⚠️ The RAG service is temporarily unavailable. Please try again in a moment.",
            author="System",
        ).send()
        return

    context        = rag_response["context_text"]
    confidence     = rag_response["confidence"]
    sources        = [dict(r) for r in rag_response["results"]]
    sources_md_str = rag_response["sources_md"]

    # Web search on low confidence — run concurrently with gap logging
    web_results: list[dict] = []
    if confidence.get("level") == "low":
        asyncio.ensure_future(_post_gap(query, confidence, sources))
        web_results = await loop.run_in_executor(None, _web_search_sync, query)

    # ── Surface similar past question ──────────────────────────────────────
    if similar:
        snippet = similar["question"][:90].rstrip()
        await cl.Message(
            content=(
                f"💡 **Similar question** from {similar['date']}:\n"
                f"> *\"{snippet}{'…' if len(similar['question']) > 90 else ''}\"*\n\n"
                f"Answering fresh below — your context may have changed."
            ),
            author="Distill", parent_id=None,
        ).send()

    # ── Build Claude messages ──────────────────────────────────────────────
    preamble = f"Context from the corpus:\n\n{context}\n\n---\n\n"

    if web_results:
        preamble += f"## Web Search Results\n\n{_format_web_context(web_results)}\n\n---\n\n"

    # Inject persistent uploaded docs
    if uploaded_docs and not attachments:
        doc_block = "\n\n---\n\n## Uploaded documents (analyze through Lenny's lens):\n"
        for doc in uploaded_docs[-MAX_UPLOADED_DOCS:]:
            doc_block += f"\n### {doc['name']}\n{doc['content'][:3000]}\n"
        preamble += doc_block

    if attachments:
        user_content = attachments_to_claude_blocks(attachments, preamble + "Question: " + query)
    else:
        user_content = preamble + "Question: " + query

    api_messages = list(history) + [{"role": "user", "content": user_content}]

    # ── Stream response ────────────────────────────────────────────────────
    client = anthropic.Anthropic()
    msg = cl.Message(content="", author="Distill", parent_id=None)
    await msg.send()

    system_prompt = await _get_system_prompt() or _build_system_prompt()
    response_text = ""
    max_retries = 2
    for attempt in range(max_retries + 1):
        try:
            with client.messages.stream(
                model=model,
                max_tokens=2048,
                system=system_prompt,
                messages=api_messages,
            ) as stream:
                for text in stream.text_stream:
                    response_text += text
                    await msg.stream_token(text)
            break
        except anthropic.APIStatusError as e:
            if attempt < max_retries and e.status_code in (429, 500, 502, 503, 529):
                await asyncio.sleep(2 ** attempt)
                continue
            await msg.update()
            await cl.Message(
                content=f"❌ API error ({e.status_code}): {e.message}. Please try again.",
                author="System",
            ).send()
            return
        except Exception as e:
            if attempt < max_retries:
                await asyncio.sleep(2 ** attempt)
                continue
            await msg.update()
            await cl.Message(
                content=f"❌ Something went wrong: {e}. Please try again.",
                author="System",
            ).send()
            return

    # ── Persist analytics metadata to DB ──────────────────────────────────
    try:
        import aiosqlite as _aio
        async with _aio.connect(DB_SYNC_URL) as _db:
            _meta = json.dumps({
                "model":       model,
                "confidence":  confidence.get("level", ""),
                "turns":       turns + 1,
                "top_score":   round(sources[0]["score"], 4) if sources else 0,
            })
            thread_id = cl.user_session.get("thread_id", "")
            await _db.execute(
                "UPDATE steps SET props = ? WHERE threadId = ? AND type = 'assistant_message' "
                "AND SUBSTR(output,1,80) = ?",
                (_meta, thread_id, (response_text or "")[:80])
            )
            await _db.commit()
    except Exception:
        pass

    # ── Append confidence + sources to the main message ───────────────────
    conf = confidence
    conf_dot = {"high": "🟢", "medium": "🟠", "low": "🔴"}[conf["level"]]
    if conf["level"] == "low" and web_results:
        conf_note = "\n\n🌐 *Limited corpus coverage — supplemented with web search.*"
    elif conf["level"] == "low":
        conf_note = "\n\n⚠️ *Limited corpus coverage here — verify key claims independently.*"
    else:
        conf_note = ""

    suffix = f"\n\n---\n{conf_dot} *{conf['label']}*{conf_note}"
    if sources:
        suffix += f"\n\n**Sources ({len(sources)})**\n\n{sources_md_str}"
    if web_results:
        suffix += f"\n\n**Web Sources ({len(web_results)})**\n\n{_format_web_sources_md(web_results)}"

    msg.content = response_text + suffix
    await msg.update()

    # ── Feedback actions ───────────────────────────────────────────────────
    msg_id = str(uuid.uuid4())
    thread_name = cl.context.session.thread_id or ""
    feedback_msg = await cl.Message(
        content="",
        author="Distill",
        actions=[
            cl.Action(name="thumbs_up",      label="👍", payload={"q": query, "a": response_text, "msg_id": msg_id}),
            cl.Action(name="thumbs_down",    label="👎", payload={"q": query, "a": response_text, "msg_id": msg_id}),
            cl.Action(name="bookmark",       label="📌", payload={"q": query, "a": response_text, "sources": sources_md_str, "thread_id": thread_id, "thread_name": thread_name}),
            cl.Action(name="copy_response",  label="📋", payload={"text": response_text}),
        ],
        parent_id=None,
    ).send()
    msg_id_map: dict = cl.user_session.get("msg_id_map", {})
    if feedback_msg:
        msg_id_map[msg_id] = feedback_msg.id
    cl.user_session.set("msg_id_map", msg_id_map)

    # ── Fire-and-forget memory extraction ─────────────────────────────────
    asyncio.ensure_future(_post_memory(query, response_text, thread_id))

    # ── Follow-up suggestions ──────────────────────────────────────────────
    suggestions = await asyncio.get_running_loop().run_in_executor(
        None, get_followup_suggestions, query, response_text,
        cl.user_session.get("user_profile", {}).get("role", "")
    )
    if suggestions:
        followup_actions = [
            cl.Action(name="followup", value=s, label=s, payload={"question": s})
            for s in suggestions
        ]
        await cl.Message(
            content="**Suggested follow-ups:**",
            author="Distill",
            actions=followup_actions,
            parent_id=None,
        ).send()

    # ── Update history ─────────────────────────────────────────────────────
    history_user_content: list | str
    if attachments:
        history_user_content = []
        for att in attachments:
            if att["kind"] == "text":
                history_user_content.append(
                    {"type": "text", "text": f"[File: {att['name']}]\n{att['content'][:1000]}"}
                )
        history_user_content.append({"type": "text", "text": query})
    else:
        history_user_content = query

    history.append({"role": "user", "content": history_user_content})
    history.append({"role": "assistant", "content": response_text})
    if len(history) > MAX_HISTORY_TURNS * 2:
        history = history[-(MAX_HISTORY_TURNS * 2):]
    cl.user_session.set("history", history)


# ── Action callbacks ──────────────────────────────────────────────────────────

@cl.action_callback("thumbs_up")
async def on_thumbs_up(action: cl.Action):
    msg_id = action.payload.get("msg_id", "")
    await save_feedback(msg_id, action.payload["q"], action.payload["a"], "up")
    msg_id_map: dict = cl.user_session.get("msg_id_map", {})
    feedback_msg_id = msg_id_map.get(msg_id)
    if feedback_msg_id:
        await cl.Message(id=feedback_msg_id, content="👍🏼  ·  👎", author="Distill", parent_id=None).update()
    else:
        await action.remove()


@cl.action_callback("thumbs_down")
async def on_thumbs_down(action: cl.Action):
    msg_id = action.payload.get("msg_id", "")
    await save_feedback(msg_id, action.payload["q"], action.payload["a"], "down")
    msg_id_map: dict = cl.user_session.get("msg_id_map", {})
    feedback_msg_id = msg_id_map.get(msg_id)
    if feedback_msg_id:
        await cl.Message(id=feedback_msg_id, content="👍  ·  👎🏼", author="Distill", parent_id=None).update()
    else:
        await action.remove()
    res = await cl.AskUserMessage(
        content="What could be better? (optional — press Enter to skip)",
        author="System",
        timeout=120,
    ).send()
    feedback_text = res["output"] if res else ""
    if feedback_text.strip():
        await save_feedback(msg_id, action.payload["q"], action.payload["a"], "down", feedback_text)



@cl.action_callback("bookmark")
async def on_bookmark(action: cl.Action):
    p = action.payload
    q, a, sources = p.get("q", ""), p.get("a", ""), p.get("sources", "")
    tid, tname    = p.get("thread_id", ""), p.get("thread_name", "")

    ok = await _post_bookmark(q, a, sources, tid, tname)
    if not ok:
        save_bookmark(q, a, sources, tid, tname)

    await cl.Message(content="📌 Saved!", author="Distill", parent_id=None).send()
    await _send_chat_settings()
    await cl.ElementSidebar.set_title("Saved Posts")
    await cl.ElementSidebar.set_elements([
        cl.Text(name=f"bookmarks_{uuid.uuid4().hex[:8]}", content=await _get_bookmarks_render() or _render_bookmarks_md(), display="inline")
    ])


@cl.action_callback("copy_response")
async def on_copy_response(action: cl.Action):
    text = action.payload.get("text", "")
    await cl.Message(
        content=f"```\n{text}\n```",
        author="Distill", parent_id=None,
    ).send()


@cl.action_callback("followup")
async def on_followup(action: cl.Action):
    question  = action.payload["question"]
    thread_id = cl.user_session.get("thread_id", "")
    asyncio.ensure_future(_post_engagement("followup_click", question, thread_id))
    await cl.Message(content=question, author="You").send()
    fake_msg = cl.Message(content=question, author="You")
    fake_msg.elements = []
    await on_message(fake_msg)


@cl.action_callback("show_saved_posts")
async def on_show_saved_posts(action: cl.Action):
    bk_md = await _get_bookmarks_render() or _render_bookmarks_md()
    await cl.ElementSidebar.set_title("📌 Saved Posts")
    await cl.ElementSidebar.set_elements([
        cl.Text(name=f"bookmarks_{uuid.uuid4().hex[:8]}", content=bk_md, display="inline")
    ])
